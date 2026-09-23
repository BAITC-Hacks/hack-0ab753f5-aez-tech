"""Adapters for the supplied 1C exports. A batch replaces its previous snapshot."""

import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
import re

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    ImportBatch, InboundLine, MonthlySale, MonthlyStock, Product, SaleLine,
    StockSnapshot, StockoutPeriod,
)


KINDS = [
    ('moq', 'MOQ / кратность'),
    ('sales_monthly', 'Помесячные продажи'),
    ('stock_monthly', 'Помесячные остатки'),
    ('transactions', 'Динамика продаж'),
    ('system_snapshot', 'Сводный файл Systeme Electric на 22.09'),
    ('inbound_iek', 'Товар в пути IEK'),
    ('stock_csv', 'Актуальные остатки CSV'),
    ('stockout_csv', 'Периоды отсутствия CSV'),
    ('catalog_csv', 'Категории и единицы CSV'),
]


def clean_key(value):
    return str(value).strip() if value is not None else ''


def decimal(value):
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value).replace(',', '.'))
    except (InvalidOperation, ValueError):
        return None


def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    for pattern in ('%Y-%m-%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            pass
    raise ValueError(f'Неверная дата: {value}')


def month_columns(header, first):
    columns = []
    for index, cell in enumerate(header[first - 1:], first):
        if not isinstance(cell, str):
            break
        match = re.search(r'20\d{2}', cell)
        if not match:
            break
        year = int(match.group())
        label = cell.lower().strip()
        names = ('янв', 'фев', 'март', 'апр', 'май', 'июн', 'июл', 'авг', 'сент', 'окт', 'нояб', 'дек')
        month = next((i for i, name in enumerate(names, 1) if label.startswith(name)), None)
        if month is None:
            raise ValueError(f'Неизвестный месяц в заголовке: {cell}')
        columns.append((index - 1, date(year, month, 1)))
    if len(columns) < 12:
        raise ValueError('Не найдено 12 месячных столбцов: проверьте формат файла.')
    return columns


def import_file(supplier, kind, filename, content, as_of=None, force=False):
    if kind not in dict(KINDS):
        raise ValueError('Неизвестный тип выгрузки.')
    if not kind.endswith('_csv') and supplier.code not in ('systeme', 'iek'):
        raise ValueError('Эти XLSX-форматы поддерживаются только для Systeme Electric и IEK.')
    digest = sha256(content).hexdigest()
    if not force and ImportBatch.objects.filter(supplier=supplier, kind=kind, sha256=digest).exists():
        return ImportBatch.objects.filter(supplier=supplier, kind=kind, sha256=digest).latest('imported_at'), False
    warnings = []
    count = 0
    with transaction.atomic():
        if kind.endswith('_csv'):
            text = content.decode('utf-8-sig')
            lines = csv.DictReader(StringIO(text), delimiter=';' if text.splitlines()[0].count(';') > text.splitlines()[0].count(',') else ',')
            if kind == 'stock_csv':
                StockSnapshot.objects.filter(product__supplier=supplier, as_of=as_of).delete()
            elif kind == 'stockout_csv':
                StockoutPeriod.objects.filter(product__supplier=supplier).delete()
            for row in lines:
                sku = clean_key(row.get('sku') or row.get('Код 1с'))
                if not sku:
                    warnings.append('Строка CSV без SKU пропущена.')
                    continue
                product, _ = Product.objects.get_or_create(supplier=supplier, sku=sku)
                warehouse = (row.get('warehouse') or row.get('Склад') or 'Все').strip()
                if kind == 'stock_csv':
                    snapshot_date = date_value(row.get('as_of') or as_of)
                    on_hand = decimal(row.get('on_hand'))
                    reserved = decimal(row.get('reserved')) or Decimal(0)
                    free = decimal(row.get('free'))
                    if on_hand is None:
                        raise ValueError(f'{sku}: нет on_hand в CSV.')
                    StockSnapshot.objects.update_or_create(product=product, warehouse=warehouse, as_of=snapshot_date, defaults={
                        'on_hand': on_hand, 'reserved': reserved, 'free': free if free is not None else on_hand - reserved,
                    })
                elif kind == 'stockout_csv':
                    starts, ends = date_value(row['starts']), date_value(row['ends'])
                    if ends < starts:
                        raise ValueError(f'{sku}: конец stockout раньше начала.')
                    StockoutPeriod.objects.create(product=product, warehouse=warehouse, starts=starts, ends=ends)
                else:
                    changed = []
                    for field in ('category', 'article', 'name', 'unit'):
                        if row.get(field):
                            setattr(product, field, row[field].strip())
                            changed.append(field)
                    for field in ('min_order', 'pack_multiple', 'purchase_to_stock_factor'):
                        if row.get(field):
                            value = decimal(row[field])
                            if value is None or value <= 0:
                                raise ValueError(f'{sku}: неверное значение {field}.')
                            setattr(product, field, value)
                            changed.append(field)
                    if 'purchase_to_stock_factor' in changed:
                        product.requires_conversion = False
                        changed.append('requires_conversion')
                    if changed:
                        product.save(update_fields=changed)
                count += 1
        else:
            wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            data = ws.iter_rows(values_only=True)
            header = next(data)
            key_col = {
                'moq': 2 if supplier.code == 'systeme' else 1,
                'sales_monthly': 1,
                'stock_monthly': 2,
                'transactions': 3,
                'inbound_iek': 0,
            }.get(kind)
            if key_col is not None and ('код' not in str(header[key_col]).lower()):
                raise ValueError(f'Неверный шаблон файла: ожидается код 1С в столбце {key_col + 1}.')
            if kind == 'transactions' and ('дата' not in str(header[0]).lower() or 'количество' not in str(header[7]).lower()):
                raise ValueError('Неверный шаблон истории продаж.')
            products = {p.sku: p for p in Product.objects.filter(supplier=supplier)}

            def product_for(sku, name='', article='', unit=''):
                sku = clean_key(sku)
                if not sku:
                    return None
                product = products.get(sku)
                if product is None:
                    product = Product.objects.create(supplier=supplier, sku=sku)
                    products[sku] = product
                changed = []
                for field, value in (('name', name), ('article', article), ('unit', unit)):
                    if value and not getattr(product, field):
                        setattr(product, field, str(value).strip())
                        changed.append(field)
                if changed:
                    product.save(update_fields=changed)
                return product

            if kind == 'moq':
                seen = set()
                Product.objects.filter(supplier=supplier).update(min_order=None, pack_multiple=None)
                for product in products.values():
                    product.min_order = None
                    product.pack_multiple = None
                if supplier.code == 'systeme':
                    next(data, None)  # Blank second row.
                for row in data:
                    sku = row[2] if supplier.code == 'systeme' else row[1]
                    product = product_for(sku, row[1] if supplier.code == 'systeme' else row[3], row[3] if supplier.code == 'systeme' else row[2])
                    if not product:
                        continue
                    if product.sku in seen:
                        warnings.append(f'{product.sku}: повторный код в файле MOQ; использована последняя строка.')
                    seen.add(product.sku)
                    value = decimal(row[4])
                    if value is None or value <= 0:
                        warnings.append(f'{product.sku}: MOQ/кратность отсутствует или ошибочна.')
                    elif supplier.code == 'systeme':
                        product.pack_multiple = value
                        product.min_order = Decimal(1)
                    else:
                        product.min_order = value
                        product.pack_multiple = Decimal(1)
                    product.save(update_fields=['min_order', 'pack_multiple'])
                    count += 1
            elif kind in ('sales_monthly', 'stock_monthly'):
                is_sale = kind == 'sales_monthly'
                first = (5 if is_sale else 5) if supplier.code == 'systeme' else (3 if is_sale else 4)
                columns = month_columns(header, first)
                if is_sale:
                    MonthlySale.objects.filter(product__supplier=supplier).delete()
                else:
                    MonthlyStock.objects.filter(product__supplier=supplier).delete()
                skip = 1 if is_sale else (2 if supplier.code == 'iek' else 2)
                for _ in range(skip):
                    next(data, None)
                buffer = []
                for row in data:
                    sku = row[1] if is_sale else (row[2] if supplier.code == 'systeme' else row[2])
                    if not sku or str(row[0]).strip() == 'Итого':
                        continue
                    name = row[0] if is_sale else (row[1] if supplier.code == 'systeme' else row[0])
                    article = row[2] if is_sale and supplier.code == 'systeme' else ''
                    unit = row[3] if not is_sale and supplier.code == 'systeme' else (row[1] if not is_sale else '')
                    product = product_for(sku, name, article, unit)
                    if is_sale and supplier.code == 'systeme':
                        pack = decimal(row[3])
                        if pack is not None and pack > 0 and product.pack_multiple is None:
                            product.pack_multiple = pack
                            product.save(update_fields=['pack_multiple'])
                    for index, month in columns:
                        raw = row[index] if index < len(row) else None
                        value = decimal(raw)
                        if is_sale:
                            buffer.append(MonthlySale(product=product, month=month, quantity=value if value is not None else 0, reported=value is not None))
                        else:
                            buffer.append(MonthlyStock(product=product, month=month, opening_quantity=value))
                    count += 1
                    if len(buffer) >= 3000:
                        (MonthlySale if is_sale else MonthlyStock).objects.bulk_create(buffer, batch_size=3000)
                        buffer.clear()
                if buffer:
                    (MonthlySale if is_sale else MonthlyStock).objects.bulk_create(buffer, batch_size=3000)
            elif kind == 'transactions':
                SaleLine.objects.filter(product__supplier=supplier).delete()
                buffer = []
                client_col = next((i for i, label in enumerate(header) if str(label).strip().lower() in ('customer_hash', 'обезличенный id клиента')), None)
                price_col = next((i for i, label in enumerate(header) if str(label).strip().lower() in ('price', 'цена')), None)
                if any(str(label).strip().lower() in ('клиент', 'фио', 'имя клиента') for label in header):
                    raise ValueError('Выгрузка содержит поле с необезличенным клиентом. Обезличьте его до импорта.')
                for row in data:
                    if not row[0] or str(row[0]).strip() == 'Итого':
                        continue
                    product = product_for(row[3], row[4], unit=row[5])
                    qty = decimal(row[7])
                    if not product or qty is None:
                        continue
                    sold_at = row[0] if isinstance(row[0], datetime) else datetime.strptime(str(row[0]), '%d.%m.%Y %H:%M:%S')
                    if timezone.is_naive(sold_at):
                        sold_at = timezone.make_aware(sold_at)
                    customer_hash = clean_key(row[client_col]) if client_col is not None and client_col < len(row) else ''
                    price = decimal(row[price_col]) if price_col is not None and price_col < len(row) else None
                    buffer.append(SaleLine(product=product, sold_at=sold_at, document=clean_key(row[1]), document_type=str(row[2]).split(' ')[0], warehouse=clean_key(row[6]), quantity=qty, customer_hash=customer_hash, unit_price=price))
                    count += 1
                    if len(buffer) >= 3000:
                        SaleLine.objects.bulk_create(buffer, batch_size=3000)
                        buffer.clear()
                if buffer:
                    SaleLine.objects.bulk_create(buffer, batch_size=3000)
            elif kind == 'system_snapshot':
                if supplier.code != 'systeme' or not as_of:
                    raise ValueError('Для сводного файла Systeme Electric требуется дата снимка.')
                snapshot_header = next(data)
                if 'код 1с' not in str(snapshot_header[2]).lower() or 'свободный остаток' not in str(snapshot_header[51]).lower():
                    raise ValueError('Неверный шаблон сводного файла Systeme Electric.')
                StockSnapshot.objects.filter(product__supplier=supplier, as_of=as_of).delete()
                InboundLine.objects.filter(product__supplier=supplier).delete()
                for row in data:
                    if not row[2] or str(row[0]).strip() in ('№', 'Итого'):
                        continue
                    product = product_for(row[2], row[3], row[1])
                    category = clean_key(row[4])
                    if product.category != category:
                        product.category = category
                        product.save(update_fields=['category'])
                    on_hand = decimal(row[49])
                    reserved = decimal(row[50]) or Decimal(0)
                    free = decimal(row[51])
                    if on_hand is not None and free is not None:
                        StockSnapshot.objects.create(product=product, warehouse='Все', as_of=as_of, on_hand=on_hand, reserved=reserved, free=free)
                    transit = decimal(row[54])
                    if transit and transit > 0:
                        InboundLine.objects.create(product=product, reference='СЭ в пути 24.09', eta=date(as_of.year, 9, 24), quantity=transit)
                    count += 1
            elif kind == 'inbound_iek':
                if supplier.code != 'iek':
                    raise ValueError('Этот формат предназначен для IEK.')
                InboundLine.objects.filter(product__supplier=supplier).delete()
                etas = []
                for text in header[3:9]:
                    matches = re.findall(r'\d{2}\.\d{2}\.\d{4}', str(text))
                    if not matches:
                        raise ValueError(f'Не найдена дата поступления в заголовке {text}')
                    etas.append(date_value(matches[-1]))
                buffer = []
                for row in data:
                    product = product_for(row[0], row[2], row[1])
                    if not product:
                        continue
                    if 'ЗАКУПАЮТСЯ БУХТАМИ' in str(row[2]).upper() and not product.requires_conversion:
                        product.requires_conversion = True
                        product.save(update_fields=['requires_conversion'])
                    for i, eta in enumerate(etas, 3):
                        qty = decimal(row[i]) if len(row) > i else None
                        if qty and qty > 0:
                            buffer.append(InboundLine(product=product, reference=str(header[i])[:160], eta=eta, quantity=qty))
                    count += 1
                InboundLine.objects.bulk_create(buffer, batch_size=2000)
            wb.close()
        batch = ImportBatch.objects.create(supplier=supplier, kind=kind, filename=filename, sha256=digest, as_of=as_of, row_count=count, warnings=warnings[:100])
        return batch, True
